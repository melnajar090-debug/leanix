import os
import shutil

import pytest
from openpyxl import load_workbook

from app import schema, store as st
from app.store import ConflictError, LockedError, Store, ValidationError


@pytest.fixture
def xlsx(tmp_path):
    path = tmp_path / "ea.xlsx"
    Store.create_new(path)
    return path


def app_values(name="CRM", **kw):
    v = {"Name": name, "Lifecycle Status": "Active"}
    v.update(kw)
    return v


def test_create_and_roundtrip(xlsx, tmp_path):
    s = Store(xlsx, tmp_path / "bk")
    s.load()
    cap = s.add("capability", {"Name": "Sales"})
    sub = s.add("capability", {"Name": "Lead Mgmt", "Parent Capability": cap})
    app = s.add("application", app_values("CRM", **{"Business Criticality": "High"}))
    itc = s.add("itcomponent", {"Name": "Postgres", "Technical Stack": "Database"})
    do = s.add("dataobject", {"Name": "Customer", "Data Classification": "Confidential"})
    bp = s.add("process", {"Name": "Order to Cash"})
    rules = {(r.source, r.target): r for r in schema.RELATION_RULES}
    s.set_links("application", app, rules[("application", "capability")], "out", [sub])
    s.set_links("application", app, rules[("application", "itcomponent")], "out", [itc])
    s.set_links("application", app, rules[("application", "dataobject")], "out", [do])
    s.set_links("process", bp, rules[("process", "capability")], "out", [cap])
    assert s.dirty
    s.save()
    assert not s.dirty

    s2 = Store(xlsx)
    s2.load()
    assert s2.get("application", app)["Business Criticality"] == "High"
    assert s2.get("capability", sub)["Parent Capability"] == cap
    assert s2.links_of("capability", sub, rules[("application", "capability")], "in") == [app]
    assert len(s2.links) == 4
    assert s2.revision == 1
    # human-readable names are written to the Relationships sheet
    ws = load_workbook(xlsx)[schema.REL_SHEET]
    assert "Lead Mgmt" in [c.value for row in ws.iter_rows() for c in row]


def test_validation(xlsx):
    s = Store(xlsx)
    s.load()
    with pytest.raises(ValidationError):
        s.add("application", {"Name": "", "Lifecycle Status": "Active"})
    with pytest.raises(ValidationError):
        s.add("application", app_values(**{"Lifecycle Status": "Bogus"}))
    s.add("application", app_values("CRM"))
    with pytest.raises(ValidationError):
        s.add("application", app_values("crm"))       # duplicate name, case-insensitive
    a = s.add("capability", {"Name": "A"})
    b = s.add("capability", {"Name": "B", "Parent Capability": a})
    with pytest.raises(ValidationError):
        s.update("capability", a, {"Name": "A", "Parent Capability": b})   # cycle


def test_delete_cascades(xlsx):
    s = Store(xlsx)
    s.load()
    cap = s.add("capability", {"Name": "Parent"})
    kid = s.add("capability", {"Name": "Kid", "Parent Capability": cap})
    app = s.add("application", app_values())
    rule = schema.RELATION_RULES[0]
    s.set_links("application", app, rule, "out", [cap])
    s.delete("capability", cap)
    assert s.get("capability", kid)["Parent Capability"] == ""
    assert not s.links


def test_two_users_disjoint_changes_merge(xlsx):
    a, b = Store(xlsx), Store(xlsx)
    a.load(); b.load()
    a.add("application", app_values("From A"))
    b.add("application", app_values("From B"))
    a.save()
    res = b.save()                                    # B's id APP-0001 collides with A's
    assert res.renumbered, "colliding ID must be renumbered, not overwritten"
    final = Store(xlsx); final.load()
    assert sorted(r["Name"] for r in final.all("application")) == ["From A", "From B"]
    assert len({r["Application ID"] for r in final.all("application")}) == 2


def test_renumbered_record_keeps_its_links(xlsx):
    a, b = Store(xlsx), Store(xlsx)
    a.load(); b.load()
    a.add("application", app_values("From A"))
    bapp = b.add("application", app_values("From B"))
    cap = b.add("capability", {"Name": "Cap"})
    b.set_links("application", bapp, schema.RELATION_RULES[0], "out", [cap])
    a.save(); b.save()
    final = Store(xlsx); final.load()
    new_id = next(r["Application ID"] for r in final.all("application") if r["Name"] == "From B")
    assert new_id != "APP-0001"
    assert final.links_of("application", new_id, schema.RELATION_RULES[0], "out") == [cap]


def test_field_level_merge_no_conflict(xlsx):
    s0 = Store(xlsx); s0.load()
    rid = s0.add("application", app_values("CRM")); s0.save()
    a, b = Store(xlsx), Store(xlsx)
    a.load(); b.load()
    a.update("application", rid, app_values("CRM", **{"Business Owner": "Alice"}))
    b.update("application", rid, app_values("CRM", **{"Technical Owner": "Bob"}))
    a.save(); b.save()
    f = Store(xlsx); f.load()
    rec = f.get("application", rid)
    assert (rec["Business Owner"], rec["Technical Owner"]) == ("Alice", "Bob")


def test_same_field_conflict_and_resolution(xlsx):
    s0 = Store(xlsx); s0.load()
    rid = s0.add("application", app_values("CRM")); s0.save()
    a, b = Store(xlsx), Store(xlsx)
    a.load(); b.load()
    a.update("application", rid, app_values("CRM", **{"Lifecycle Status": "Phase Out"}))
    b.update("application", rid, app_values("CRM", **{"Lifecycle Status": "End of Life"}))
    a.save()
    with pytest.raises(ConflictError) as ei:
        b.save()
    c = ei.value.conflicts[0]
    assert (c.field, c.mine, c.theirs) == ("Lifecycle Status", "End of Life", "Phase Out")
    chk = Store(xlsx); chk.load()
    assert chk.get("application", rid)["Lifecycle Status"] == "Phase Out"   # nothing was lost/written
    b.save(resolve="mine")
    chk.load()
    assert chk.get("application", rid)["Lifecycle Status"] == "End of Life"


def test_conflict_resolution_theirs(xlsx):
    s0 = Store(xlsx); s0.load()
    rid = s0.add("application", app_values("CRM")); s0.save()
    a, b = Store(xlsx), Store(xlsx)
    a.load(); b.load()
    a.update("application", rid, app_values("CRM", **{"Lifecycle Status": "Phase Out"}))
    b.update("application", rid, app_values("CRM", **{"Lifecycle Status": "End of Life"}))
    a.save(); b.save(resolve="theirs")
    assert b.get("application", rid)["Lifecycle Status"] == "Phase Out"
    assert not b.dirty


def test_edit_vs_remote_delete(xlsx):
    s0 = Store(xlsx); s0.load()
    rid = s0.add("application", app_values("CRM")); s0.save()
    a, b = Store(xlsx), Store(xlsx)
    a.load(); b.load()
    a.delete("application", rid); a.save()
    b.update("application", rid, app_values("CRM", **{"Business Owner": "Bob"}))
    with pytest.raises(ConflictError) as ei:
        b.save()
    assert ei.value.conflicts[0].kind == "deleted_remote"


def test_remote_delete_of_link_target_drops_link(xlsx):
    s0 = Store(xlsx); s0.load()
    cap = s0.add("capability", {"Name": "Cap"}); app = s0.add("application", app_values()); s0.save()
    a, b = Store(xlsx), Store(xlsx)
    a.load(); b.load()
    a.delete("capability", cap); a.save()
    b.set_links("application", app, schema.RELATION_RULES[0], "out", [cap])
    res = b.save()
    assert any("relationship" in w for w in res.warnings)
    assert not b.links


def test_locked_file_raises_and_keeps_changes(xlsx, monkeypatch):
    s = Store(xlsx); s.load()
    s.add("application", app_values())
    real = os.replace

    def locked(*a, **k):
        raise PermissionError("locked by Excel")
    monkeypatch.setattr(st.os, "replace", locked)
    monkeypatch.setattr(st.time, "sleep", lambda *_: None)
    with pytest.raises(LockedError):
        s.save()
    assert s.dirty                                       # changes still pending
    assert not list(xlsx.parent.glob(".leanixlite-*"))   # no temp file left behind
    monkeypatch.setattr(st.os, "replace", real)
    s.save()                                             # retry succeeds
    assert not s.dirty


def test_excel_owner_file_detected(xlsx):
    s = Store(xlsx); s.load()
    s.add("application", app_values())
    (xlsx.parent / "~$ea.xlsx").write_text("x")
    with pytest.raises(LockedError):
        s.save()
    (xlsx.parent / "~$ea.xlsx").unlink()
    s.save()


def test_backup_created_and_pruned(xlsx, tmp_path):
    s = Store(xlsx, tmp_path / "bk"); s.load()
    s.BACKUPS_KEPT = 2
    for i in range(4):
        s.add("application", app_values(f"App{i}"))
        s.save()
    assert 1 <= len(list((tmp_path / "bk").glob("*.xlsx"))) <= 2


def test_rows_added_in_excel_get_ids_and_extra_columns_survive(xlsx):
    wb = load_workbook(xlsx)
    ws = wb["Applications"]
    ws.cell(row=1, column=ws.max_column + 1, value="My Notes")
    ws.append([None, "Typed In Excel", None, "Active"])
    wb.save(xlsx)
    s = Store(xlsx); s.load()
    assert any("row without ID" in w for w in s.load_warnings)
    rec = next(r for r in s.all("application") if r["Name"] == "Typed In Excel")
    assert rec["Application ID"].startswith("APP-")
    assert s.dirty                                       # needs a save to persist the new ID
    s.save()
    assert "My Notes" in [c.value for c in load_workbook(xlsx)["Applications"][1]]


def test_disk_changed_detection_and_conflicted_copies(xlsx):
    s = Store(xlsx); s.load()
    assert not s.disk_changed()
    other = Store(xlsx); other.load(); other.add("application", app_values()); other.save()
    assert s.disk_changed()
    shutil.copy(xlsx, xlsx.parent / "ea-LAPTOP.xlsx")
    assert st.find_conflicted_copies(xlsx) == ["ea-LAPTOP.xlsx"]


def test_open_or_create_in_folder(tmp_path, monkeypatch):
    from app import config, ui_main
    monkeypatch.setattr(config, "BACKUP_DIR", tmp_path / "bk")
    folder = tmp_path / "OneDrive" / "LeanIX"
    folder.mkdir(parents=True)
    s1 = ui_main.open_or_create("folder", str(folder))          # empty folder -> creates database
    assert (folder / config.DB_FILENAME).exists()
    s1.add("application", app_values("Kept")); s1.save()
    s2 = ui_main.open_or_create("folder", str(folder))          # existing -> opens it, data intact
    assert [r["Name"] for r in s2.all("application")] == ["Kept"]
