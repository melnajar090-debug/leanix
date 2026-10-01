"""Excel-backed data store with record-level merge-on-save.

Why not just overwrite the workbook?  OneDrive syncs whole files, so two people
saving from stale copies would silently destroy each other's work.  Instead the
store remembers the state it loaded (the *base*), and on save it:

  1. re-reads the file as it is on disk right now (*theirs*),
  2. works out what *you* changed relative to the base (added / edited / deleted
     records and links),
  3. applies only those changes on top of theirs (three-way, field-level merge),
  4. reports real conflicts (same field edited differently, edit vs delete),
  5. writes to a temp file and atomically replaces the workbook.

A file that is open in Excel (or briefly locked by the OneDrive sync client) raises
``LockedError``; the app catches it and lets the user retry.
"""
from __future__ import annotations

import copy
import getpass
import os
import re
import shutil
import tempfile
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, NamedTuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import schema
from .schema import ENTITIES, ENTITY_ORDER, REL_COLUMNS, REL_SHEET, META_SHEET, AUDIT_FIELDS

Record = dict[str, str]


# --------------------------------------------------------------------------- errors
class StoreError(Exception):
    pass


class LockedError(StoreError):
    """The workbook cannot be written (open in Excel, or locked by the sync client)."""


class ConflictError(StoreError):
    def __init__(self, conflicts: list["Conflict"]):
        super().__init__(f"{len(conflicts)} conflicting change(s)")
        self.conflicts = conflicts


class ValidationError(StoreError):
    def __init__(self, errors: list[str]):
        super().__init__("; ".join(errors))
        self.errors = errors


@dataclass
class Conflict:
    entity: str
    record_id: str
    name: str
    kind: str        # "field" | "deleted_remote" | "deleted_local"
    field: str = ""
    mine: str = ""
    theirs: str = ""

    def describe(self) -> str:
        label = f"{ENTITIES[self.entity].label} {self.record_id} '{self.name}'"
        if self.kind == "field":
            return (f"{label}: field '{self.field}' - you set '{self.mine}', "
                    f"someone else set '{self.theirs}'")
        if self.kind == "deleted_remote":
            return f"{label}: you edited it, but someone else deleted it"
        return f"{label}: you deleted it, but someone else edited it"


class Link(NamedTuple):
    src_type: str
    src_id: str
    tgt_type: str
    tgt_id: str
    rel: str


@dataclass
class Pending:
    added: list[tuple[str, str]] = field(default_factory=list)
    modified: list[tuple[str, str]] = field(default_factory=list)
    deleted: list[tuple[str, str]] = field(default_factory=list)
    links_added: set[Link] = field(default_factory=set)
    links_removed: set[Link] = field(default_factory=set)

    def __bool__(self) -> bool:
        return bool(self.added or self.modified or self.deleted
                    or self.links_added or self.links_removed)

    @property
    def count(self) -> int:
        return (len(self.added) + len(self.modified) + len(self.deleted)
                + len(self.links_added) + len(self.links_removed))


@dataclass
class SaveResult:
    warnings: list[str] = field(default_factory=list)
    renumbered: dict[str, str] = field(default_factory=dict)
    revision: int = 0


@dataclass
class _Snapshot:
    records: dict[str, dict[str, Record]]
    links: dict[Link, dict[str, str]]
    extra_cols: dict[str, list[str]]
    revision: int = 0
    warnings: list[str] = field(default_factory=list)
    fresh: list[tuple[str, str]] = field(default_factory=list)   # rows that got an ID on load


# --------------------------------------------------------------------------- helpers
def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def current_user() -> str:
    try:
        return getpass.getuser()
    except Exception:  # pragma: no cover - exotic environments
        return os.environ.get("USERNAME", "unknown")


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _id_number(prefix: str, rid: str) -> int:
    m = re.fullmatch(re.escape(prefix) + r"-(\d+)", rid or "")
    return int(m.group(1)) if m else 0


def next_id(prefix: str, existing: Iterable[str]) -> str:
    return f"{prefix}-{max([_id_number(prefix, i) for i in existing] + [0]) + 1:04d}"


def excel_owner_file_present(path: Path) -> bool:
    """True if Excel's ``~$name.xlsx`` owner file shows the workbook is open in Excel."""
    try:
        for entry in os.listdir(path.parent):
            if entry.startswith("~$") and len(entry) > 2 and path.name.endswith(entry[2:]):
                return True
    except OSError:
        pass
    return False


def find_conflicted_copies(path: Path) -> list[str]:
    """Heuristic: OneDrive-style duplicates such as ``name-COMPUTER.xlsx`` next to the workbook."""
    stem, suffix = path.stem, path.suffix
    found = []
    try:
        for entry in os.listdir(path.parent):
            if entry == path.name or entry.startswith("~$"):
                continue
            if entry.startswith(stem + "-") and entry.endswith(suffix):
                found.append(entry)
    except OSError:
        pass
    return sorted(found)


# --------------------------------------------------------------------------- reading
def _read_workbook(path: Path) -> _Snapshot:
    last_err: Exception | None = None
    wb = None
    for _ in range(3):
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            break
        except PermissionError as e:   # briefly locked by sync client
            last_err = e
            time.sleep(0.4)
    if wb is None:
        raise LockedError(f"Cannot read '{path.name}': {last_err}")

    snap = _Snapshot({ek: {} for ek in ENTITY_ORDER}, {}, {})
    try:
        names = set(wb.sheetnames)

        def rows(sheet: str):
            if sheet not in names:
                return [], []
            it = wb[sheet].iter_rows(values_only=True)
            try:
                header = [_cell_to_str(h) for h in next(it)]
            except StopIteration:
                return [], []
            body = [[_cell_to_str(c) for c in r] for r in it]
            return header, [r for r in body if any(r)]

        for ek in ENTITY_ORDER:
            ent = ENTITIES[ek]
            header, body = rows(ent.sheet)
            known = set(ent.columns)
            snap.extra_cols[ek] = [h for h in header if h and h not in known]
            idx = {h: i for i, h in enumerate(header) if h}
            for r in body:
                rec: Record = {c: "" for c in ent.columns}
                for h, i in idx.items():
                    if i < len(r):
                        rec[h] = r[i]
                rid = rec[ent.id_field]
                if not rid:
                    rid = next_id(ent.prefix, list(snap.records[ek]))
                    rec[ent.id_field] = rid
                    rec["Created By"] = rec["Created By"] or current_user()
                    rec["Created At"] = rec["Created At"] or now_str()
                    snap.fresh.append((ek, rid))
                    snap.warnings.append(
                        f"{ent.sheet}: row without ID received {rid} (pending until you save)")
                if rid in snap.records[ek]:
                    snap.warnings.append(f"{ent.sheet}: duplicate ID {rid} ignored")
                    continue
                snap.records[ek][rid] = rec

        header, body = rows(REL_SHEET)
        idx = {h: i for i, h in enumerate(header) if h}
        need = ["Source Type", "Source ID", "Target Type", "Target ID"]
        if body and not all(n in idx for n in need):
            snap.warnings.append(f"{REL_SHEET}: required columns missing, relationships ignored")
            body = []
        for r in body:
            def g(col):
                i = idx.get(col)
                return r[i] if i is not None and i < len(r) else ""
            link = Link(g("Source Type"), g("Source ID"), g("Target Type"), g("Target ID"),
                        g("Relation Type"))
            if link.src_type in ENTITIES and link.tgt_type in ENTITIES:
                snap.links[link] = {"Created By": g("Created By"), "Created At": g("Created At")}

        if META_SHEET in names:
            for r in wb[META_SHEET].iter_rows(values_only=True):
                if r and r[0] == "Revision":
                    try:
                        snap.revision = int(r[1])
                    except (TypeError, ValueError):
                        pass
    finally:
        wb.close()
    return snap


# --------------------------------------------------------------------------- writing
_HEADER_FILL = PatternFill("solid", start_color="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_sheet(ws, headers: list[str], n_rows: int, widths: dict[str, int] | None = None):
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        width = (widths or {}).get(headers[c - 1], max(14, min(40, len(headers[c - 1]) + 4)))
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(n_rows + 1, 2)}"


def _build_workbook(records, links, extra_cols, revision: int) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for ek in ENTITY_ORDER:
        ent = ENTITIES[ek]
        headers = ent.columns + list(extra_cols.get(ek, []))
        ws = wb.create_sheet(ent.sheet)
        ws.append(headers)
        recs = list(records[ek].values())
        for rec in recs:
            ws.append([rec.get(h, "") for h in headers])
        widths = {f.name: (60 if f.kind == "long" else 28) for f in ent.fields}
        widths[ent.id_field] = 12
        _style_sheet(ws, headers, len(recs), widths)
        last = max(len(recs) + 200, 500)   # keep dropdowns usable for rows added in Excel
        for f in ent.fields:
            if f.kind == "choice":
                col = get_column_letter(headers.index(f.name) + 1)
                dv = DataValidation(type="list", formula1='"' + ",".join(f.choices) + '"',
                                    allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{col}2:{col}{last}")

    ws = wb.create_sheet(REL_SHEET)
    ws.append(REL_COLUMNS)
    n = 0
    for link, meta in links.items():
        s = records.get(link.src_type, {}).get(link.src_id, {})
        t = records.get(link.tgt_type, {}).get(link.tgt_id, {})
        ws.append([link.src_type, link.src_id, s.get("Name", ""), link.rel,
                   link.tgt_type, link.tgt_id, t.get("Name", ""),
                   meta.get("Created By", ""), meta.get("Created At", "")])
        n += 1
    _style_sheet(ws, REL_COLUMNS, n,
                 {"Source Name": 30, "Target Name": 30, "Source Type": 14, "Target Type": 14})

    ws = wb.create_sheet(META_SHEET)
    for row in [("Key", "Value"),
                ("Schema Version", schema.SCHEMA_VERSION),
                ("Revision", revision),
                ("Last Saved By", current_user()),
                ("Last Saved At", now_str()),
                ("Note", "Managed by LeanIX Lite. Please edit through the app, not directly in Excel.")]:
        ws.append(list(row))
    _style_sheet(ws, ["Key", "Value"], 5, {"Key": 18, "Value": 60})
    return wb


def _atomic_write(path: Path, wb: Workbook, retries: int = 3, delay: float = 0.5):
    fd, tmp = tempfile.mkstemp(prefix=".leanixlite-", suffix=".tmp", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        last: Exception | None = None
        for attempt in range(retries):
            try:
                os.replace(tmp, path)
                return
            except PermissionError as e:
                last = e
                if attempt < retries - 1:
                    time.sleep(delay)
        raise LockedError(
            f"'{path.name}' is locked - it is probably open in Excel or being synced by OneDrive.") from last
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


# --------------------------------------------------------------------------- the store
class Store:
    BACKUPS_KEPT = 15

    def __init__(self, path: str | os.PathLike, backup_dir: str | os.PathLike | None = None):
        self.path = Path(path)
        self.backup_dir = Path(backup_dir) if backup_dir else None
        self.records: dict[str, dict[str, Record]] = {ek: {} for ek in ENTITY_ORDER}
        self.links: dict[Link, dict[str, str]] = {}
        self.extra_cols: dict[str, list[str]] = {ek: [] for ek in ENTITY_ORDER}
        self.revision = 0
        self.load_warnings: list[str] = []
        self._base_records = copy.deepcopy(self.records)
        self._base_links: dict[Link, dict[str, str]] = {}
        self._disk_stamp: tuple[int, int] | None = None

    # ----- creation / loading
    @classmethod
    def create_new(cls, path, backup_dir=None) -> "Store":
        path = Path(path)
        if path.exists():
            raise StoreError(f"'{path}' already exists")
        st = cls(path, backup_dir)
        wb = _build_workbook(st.records, {}, st.extra_cols, 0)
        _atomic_write(path, wb)
        st.load()
        return st

    def load(self):
        """Replace in-memory state with what is on disk (discards unsaved local changes)."""
        snap = _read_workbook(self.path)
        self.records = snap.records
        self.links = snap.links
        self.extra_cols = snap.extra_cols
        self.revision = snap.revision
        self.load_warnings = snap.warnings
        # Rows that had no ID were numbered locally: they are pending additions, not base.
        self._base_records = copy.deepcopy(snap.records)
        for ek, rid in snap.fresh:
            self._base_records[ek].pop(rid, None)
        self._base_links = copy.deepcopy(snap.links)
        self._disk_stamp = self._stamp()

    def _stamp(self) -> tuple[int, int] | None:
        try:
            st = os.stat(self.path)
            return (st.st_mtime_ns, st.st_size)
        except OSError:
            return None

    def disk_changed(self) -> bool:
        """True if the file changed on disk since we last loaded/saved it (e.g. OneDrive sync)."""
        return self._stamp() != self._disk_stamp

    # ----- reading
    def get(self, ek: str, rid: str) -> Record | None:
        return self.records[ek].get(rid)

    def all(self, ek: str) -> list[Record]:
        return list(self.records[ek].values())

    def name_of(self, ek: str, rid: str) -> str:
        rec = self.records[ek].get(rid)
        return rec.get("Name", "") if rec else ""

    def label_of(self, ek: str, rid: str) -> str:
        if not rid:
            return ""
        rec = self.records[ek].get(rid)
        return f"{rid} · {rec['Name']}" if rec else f"{rid} (missing)"

    def links_of(self, ek: str, rid: str, rule: schema.RelationRule, direction: str) -> list[str]:
        """IDs linked to a record under a rule. direction 'out' = record is source."""
        out = []
        for l in self.links:
            if l.rel != rule.label:
                continue
            if direction == "out" and (l.src_type, l.src_id, l.tgt_type) == (ek, rid, rule.target) \
                    and l.src_type == rule.source:
                out.append(l.tgt_id)
            elif direction == "in" and (l.tgt_type, l.tgt_id, l.src_type) == (ek, rid, rule.source) \
                    and l.tgt_type == rule.target:
                out.append(l.src_id)
        return out

    def related(self, ek: str, rid: str) -> list[tuple[str, str, str]]:
        """Every neighbour as (rel wording, entity key, id), regardless of rule."""
        res = []
        for l in self.links:
            if (l.src_type, l.src_id) == (ek, rid):
                res.append((l.rel, l.tgt_type, l.tgt_id))
            elif (l.tgt_type, l.tgt_id) == (ek, rid):
                res.append((f"{l.rel} (inverse)", l.src_type, l.src_id))
        return res

    # ----- pending changes
    def pending(self) -> Pending:
        p = Pending()
        for ek in ENTITY_ORDER:
            mine, base = self.records[ek], self._base_records[ek]
            p.added += [(ek, i) for i in mine if i not in base]
            p.deleted += [(ek, i) for i in base if i not in mine]
            p.modified += [(ek, i) for i in mine if i in base and mine[i] != base[i]]
        p.links_added = set(self.links) - set(self._base_links)
        p.links_removed = set(self._base_links) - set(self.links)
        return p

    @property
    def dirty(self) -> bool:
        return bool(self.pending())

    # ----- validation
    def validate(self, ek: str, values: Record, rid: str | None = None) -> list[str]:
        ent = ENTITIES[ek]
        errs = []
        for f in ent.fields:
            v = values.get(f.name, "").strip()
            if f.required and not v:
                errs.append(f"'{f.name}' is required.")
            if f.kind == "choice" and v and v not in f.choices:
                errs.append(f"'{f.name}' must be one of: {', '.join(f.choices)}.")
            if f.kind == "ref" and v and v not in self.records[f.ref]:
                errs.append(f"'{f.name}': {v} does not exist.")
        name = values.get("Name", "").strip().lower()
        if name:
            for other_id, rec in self.records[ek].items():
                if other_id != rid and rec["Name"].strip().lower() == name:
                    errs.append(f"A {ent.label.lower()} named '{values['Name'].strip()}' already exists ({other_id}).")
                    break
        if ek == "capability" and rid and values.get("Parent Capability"):
            cur, seen = values["Parent Capability"], {rid}
            while cur:
                if cur in seen:
                    errs.append("Parent Capability would create a cycle.")
                    break
                seen.add(cur)
                cur = self.records["capability"].get(cur, {}).get("Parent Capability", "")
        return errs

    # ----- mutation
    def add(self, ek: str, values: Record) -> str:
        errs = self.validate(ek, values)
        if errs:
            raise ValidationError(errs)
        ent = ENTITIES[ek]
        rid = next_id(ent.prefix, self.records[ek])
        user, ts = current_user(), now_str()
        rec = {c: "" for c in ent.columns}
        rec.update({f.name: values.get(f.name, "").strip() for f in ent.fields})
        rec.update({ent.id_field: rid, "Created By": user, "Created At": ts,
                    "Modified By": user, "Modified At": ts})
        for extra in self.extra_cols.get(ek, []):
            rec[extra] = ""
        self.records[ek][rid] = rec
        return rid

    def update(self, ek: str, rid: str, values: Record):
        errs = self.validate(ek, values, rid)
        if errs:
            raise ValidationError(errs)
        rec = self.records[ek][rid]
        new = dict(rec)
        for f in ENTITIES[ek].fields:
            new[f.name] = values.get(f.name, "").strip()
        if new != rec:
            new["Modified By"], new["Modified At"] = current_user(), now_str()
            self.records[ek][rid] = new

    def delete(self, ek: str, rid: str) -> list[str]:
        """Delete a record and its links; clears dangling parent references. Returns notes."""
        notes = []
        if rid not in self.records[ek]:
            return notes
        del self.records[ek][rid]
        for l in [l for l in self.links
                  if (l.src_type, l.src_id) == (ek, rid) or (l.tgt_type, l.tgt_id) == (ek, rid)]:
            del self.links[l]
        for oek in ENTITY_ORDER:
            for f in ENTITIES[oek].fields:
                if f.kind == "ref" and f.ref == ek:
                    for orid, rec in self.records[oek].items():
                        if rec[f.name] == rid:
                            rec[f.name] = ""
                            rec["Modified By"], rec["Modified At"] = current_user(), now_str()
                            notes.append(f"Cleared '{f.name}' on {orid}")
        return notes

    def set_links(self, ek: str, rid: str, rule: schema.RelationRule, direction: str,
                  other_ids: Iterable[str]):
        """Make the set of links of one rule/direction for a record exactly ``other_ids``."""
        other_ids = set(other_ids)
        other_ek = rule.target if direction == "out" else rule.source
        current = set(self.links_of(ek, rid, rule, direction))
        user, ts = current_user(), now_str()
        for oid in other_ids - current:
            if oid not in self.records[other_ek]:
                continue
            link = (Link(ek, rid, other_ek, oid, rule.label) if direction == "out"
                    else Link(other_ek, oid, ek, rid, rule.label))
            self.links[link] = {"Created By": user, "Created At": ts}
        for oid in current - other_ids:
            link = (Link(ek, rid, other_ek, oid, rule.label) if direction == "out"
                    else Link(other_ek, oid, ek, rid, rule.label))
            self.links.pop(link, None)

    # ----- saving
    def save(self, resolve: str | None = None) -> SaveResult:
        """Merge local changes into the file on disk and write it.

        ``resolve``: None -> raise ConflictError on conflicts; "mine" -> my values win;
        "theirs" -> the other person's values win.
        """
        if excel_owner_file_present(self.path):
            raise LockedError(f"'{self.path.name}' appears to be open in Excel "
                              f"(found its ~$ owner file). Close it in Excel and retry.")
        theirs = _read_workbook(self.path)
        merged_records, merged_links, result, conflicts = self._merge(theirs, resolve)
        if conflicts and resolve is None:
            raise ConflictError(conflicts)

        wb = _build_workbook(merged_records, merged_links, theirs.extra_cols,
                             theirs.revision + 1)
        self._backup()
        _atomic_write(self.path, wb)

        result.revision = theirs.revision + 1
        self.records = merged_records
        self.links = merged_links
        self.extra_cols = theirs.extra_cols
        self.revision = result.revision
        self._base_records = copy.deepcopy(merged_records)
        self._base_links = copy.deepcopy(merged_links)
        self._disk_stamp = self._stamp()
        self.load_warnings = []
        return result

    def _backup(self):
        if not self.backup_dir or not self.path.exists():
            return
        try:
            self.backup_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
            shutil.copy2(self.path, self.backup_dir / f"{self.path.stem}-{stamp}{self.path.suffix}")
            old = sorted(self.backup_dir.glob(f"{self.path.stem}-*{self.path.suffix}"))
            for f in old[:-self.BACKUPS_KEPT]:
                f.unlink(missing_ok=True)
        except OSError:
            pass   # a failed backup must never block saving

    def _merge(self, theirs: _Snapshot, resolve: str | None):
        result = SaveResult()
        conflicts: list[Conflict] = []
        out_rec = copy.deepcopy(theirs.records)
        renum: dict[tuple[str, str], str] = {}
        p = self.pending()
        audit = set(AUDIT_FIELDS)

        # additions (allocate a fresh ID if someone else took ours)
        for ek, rid in p.added:
            mine = copy.deepcopy(self.records[ek][rid])
            ent = ENTITIES[ek]
            new_id = rid
            if rid in out_rec[ek]:
                new_id = next_id(ent.prefix, list(out_rec[ek]) + list(self.records[ek]))
                renum[(ek, rid)] = new_id
                result.renumbered[f"{ek}:{rid}"] = new_id
                result.warnings.append(
                    f"{ent.label} {rid} '{mine['Name']}' was renumbered to {new_id} "
                    f"because someone else created {rid} in the meantime.")
            mine[ent.id_field] = new_id
            out_rec[ek][new_id] = mine
        # remap references from my added/edited records to renumbered ones
        for (ek, old), new in renum.items():
            for oek in ENTITY_ORDER:
                for f in ENTITIES[oek].fields:
                    if f.kind == "ref" and f.ref == ek:
                        for orid, rec in out_rec[oek].items():
                            if rec[f.name] == old and (oek, orid) in {*p.added, *p.modified}:
                                rec[f.name] = new

        # edits
        for ek, rid in p.modified:
            mine, base = self.records[ek][rid], self._base_records[ek][rid]
            dest = out_rec[ek].get(rid)
            name = mine.get("Name", "")
            if dest is None:
                conflicts.append(Conflict(ek, rid, name, "deleted_remote"))
                if resolve == "mine":
                    out_rec[ek][rid] = copy.deepcopy(mine)
                continue
            changed = [c for c in mine if c not in audit and mine.get(c, "") != base.get(c, "")]
            applied = False
            for c in changed:
                cur, orig, want = dest.get(c, ""), base.get(c, ""), mine.get(c, "")
                if cur != orig and cur != want:
                    conflicts.append(Conflict(ek, rid, name, "field", c, want, cur))
                    if resolve != "mine":
                        continue
                dest[c] = want
                applied = True
            if applied:
                for c in AUDIT_FIELDS:
                    if c.startswith("Modified"):
                        dest[c] = mine.get(c, "")

        # deletions
        for ek, rid in p.deleted:
            base = self._base_records[ek][rid]
            dest = out_rec[ek].get(rid)
            if dest is None:
                continue
            if dest != base:
                conflicts.append(Conflict(ek, rid, base.get("Name", ""), "deleted_local"))
                if resolve != "mine":
                    continue
            del out_rec[ek][rid]

        # links: disk + my additions - my removals, with renumbered IDs
        def remap(l: Link) -> Link:
            return Link(l.src_type, renum.get((l.src_type, l.src_id), l.src_id),
                        l.tgt_type, renum.get((l.tgt_type, l.tgt_id), l.tgt_id), l.rel)

        out_links = copy.deepcopy(theirs.links)
        for l in p.links_removed:
            out_links.pop(l, None)
        for l in p.links_added:
            out_links[remap(l)] = copy.deepcopy(self.links[l])

        # referential integrity after the merge
        dropped = 0
        for l in list(out_links):
            if l.src_id not in out_rec[l.src_type] or l.tgt_id not in out_rec[l.tgt_type]:
                del out_links[l]
                dropped += 1
        if dropped:
            result.warnings.append(f"{dropped} relationship(s) pointed to records deleted by "
                                   f"someone else and were removed.")
        for oek in ENTITY_ORDER:
            for f in ENTITIES[oek].fields:
                if f.kind == "ref":
                    for orid, rec in out_rec[oek].items():
                        if rec[f.name] and rec[f.name] not in out_rec[f.ref]:
                            result.warnings.append(
                                f"{orid}: '{f.name}' pointed to missing {rec[f.name]}; cleared.")
                            rec[f.name] = ""
        return out_rec, out_links, result, conflicts
